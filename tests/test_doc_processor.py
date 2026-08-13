"""
Unit Tests for Document Transliteration Processor (.txt, .docx, .pdf)
"""

import io
import pytest
from app.core.doc_processor import doc_processor

def test_txt_transliteration():
    text_content = "O'zbekiston Respublikasi kelajagi buyuk davlatdir."
    file_bytes = text_content.encode("utf-8")

    out_bytes, out_name, mime = doc_processor.process_document(
        file_bytes=file_bytes,
        filename="test_doc.txt",
        direction="latin_to_cyrillic"
    )

    result_text = out_bytes.decode("utf-8")
    assert result_text == "Ўзбекистон Республикаси келажаги буюк давлатдир."
    assert out_name == "test_doc_transliterated.txt"
    assert mime == "text/plain"

def test_docx_transliteration():
    import docx

    doc = docx.Document()
    doc.add_heading("O'zbekiston", level=1)
    p = doc.add_paragraph("Respublikasi kelajagi buyuk davlatdir.")
    p.add_run(" Ma'no va e'tibor.").bold = True

    in_io = io.BytesIO()
    doc.save(in_io)
    file_bytes = in_io.getvalue()

    out_bytes, out_name, mime = doc_processor.process_document(
        file_bytes=file_bytes,
        filename="hujjat.docx",
        direction="latin_to_cyrillic"
    )

    assert out_name == "hujjat_transliterated.docx"
    assert mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    # Verify content inside output docx
    res_doc = docx.Document(io.BytesIO(out_bytes))
    full_text = " ".join([p.text for p in res_doc.paragraphs])
    assert "Ўзбекистон" in full_text
    assert "Республикаси" in full_text
    assert "Маъно" in full_text

def test_pdf_transliteration():
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet

    in_io = io.BytesIO()
    pdf_doc = SimpleDocTemplate(in_io, pagesize=letter)
    styles = getSampleStyleSheet()
    story = [Paragraph("O'zbekiston Respublikasi", styles["Normal"])]
    pdf_doc.build(story)
    file_bytes = in_io.getvalue()

    out_bytes, out_name, mime = doc_processor.process_document(
        file_bytes=file_bytes,
        filename="hujjat.pdf",
        direction="latin_to_cyrillic"
    )

    assert out_name == "hujjat_transliterated.pdf"
    assert mime == "application/pdf"
    assert len(out_bytes) > 0
